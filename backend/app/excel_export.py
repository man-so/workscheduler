from __future__ import annotations

from calendar import monthrange
from collections import Counter, defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import models, schemas


CATEGORY_COLORS = {
    "DAY": "D9EAD3",
    "EARLY_DAY": "D0E0E3",
    "NIGHT": "D9D2E9",
    "EARLY_NIGHT": "CFE2F3",
    "OFF": "F3F4F6",
    "LEAVE": "FCE5CD",
    "OTHER": "FFF2CC",
    "REGULAR": "E5E7EB",
    "ANNUAL": "FCE5CD",
    "COMPENSATORY": "CFE2F3",
    "SPECIAL": "D9D2E9",
    "UNSPECIFIED": "F3F4F6",
}

CATEGORY_LABELS = {
    "DAY": "주간",
    "EARLY_DAY": "조기",
    "NIGHT": "야간",
    "EARLY_NIGHT": "야간 조기",
    "OFF": "휴무",
    "LEAVE": "연차",
    "OTHER": "기타",
    "REGULAR": "정기휴무",
    "ANNUAL": "연차",
    "COMPENSATORY": "대체휴무",
    "SPECIAL": "특별휴무",
    "UNSPECIFIED": "유형 미지정",
}


def _label(value: str | None) -> str:
    if not value:
        return "기타"
    return CATEGORY_LABELS.get(value, value)


def _fill(category: str | None) -> PatternFill:
    color = CATEGORY_COLORS.get(category or "OTHER", CATEGORY_COLORS["OTHER"])
    return PatternFill("solid", fgColor=color)


def _header(ws, title: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws.freeze_panes = "A3"


def _apply_table_style(ws, max_row: int, max_column: int) -> None:
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")


def _assignment_label(item: schemas.CalendarAssignment) -> str:
    if item.status == models.AssignmentStatus.WORK:
        return item.shift_type_name or _label(item.category)
    return item.leave_label or _label(item.leave_type_code or item.status)


def _ordered_groups(day: schemas.CalendarDay) -> list[tuple[str, list[schemas.CalendarAssignment]]]:
    order = ["DAY", "EARLY_DAY", "NIGHT", "EARLY_NIGHT", "OTHER", "ANNUAL", "COMPENSATORY", "REGULAR", "SPECIAL", "UNSPECIFIED", "LEAVE", "OFF"]
    return sorted(day.groups.items(), key=lambda item: order.index(item[0]) if item[0] in order else len(order))


def build_schedule_workbook(
    team: models.Team,
    version: models.ScheduleVersion,
    calendar: schemas.MonthlyScheduleRead,
    validation: schemas.ScheduleValidationRead,
) -> BytesIO:
    wb = Workbook()
    calendar_ws = wb.active
    calendar_ws.title = "월간 달력"
    _build_calendar_sheet(calendar_ws, team, version, calendar)
    _build_employee_sheet(wb.create_sheet("직원별 현황"), calendar)
    _build_shift_count_sheet(wb.create_sheet("근무유형 통계"), calendar)
    _build_fairness_sheet(wb.create_sheet("공정성 통계"), calendar)
    _build_validation_sheet(wb.create_sheet("검증 결과"), validation)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_calendar_sheet(ws, team: models.Team, version: models.ScheduleVersion, calendar: schemas.MonthlyScheduleRead) -> None:
    title = f"{team.name} {calendar.year}년 {calendar.month}월 근무표"
    _header(ws, title)
    ws["A2"] = f"버전 {version.version_no} / 상태 {version.status}"
    ws["A2"].font = Font(bold=True)
    weekdays = ["월", "화", "수", "목", "금", "토", "일"]
    for index, weekday in enumerate(weekdays, start=1):
        cell = ws.cell(row=4, column=index, value=weekday)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(index)].width = 25

    first_day = date(calendar.year, calendar.month, 1)
    day_by_date = {date.fromisoformat(day.local_date.isoformat()): day for day in calendar.days}
    row = 5
    column = first_day.weekday() + 1
    for day_number in range(1, monthrange(calendar.year, calendar.month)[1] + 1):
        current = date(calendar.year, calendar.month, day_number)
        day = day_by_date[current]
        lines = [f"{day_number}일"]
        for group, assignments in _ordered_groups(day):
            names = ", ".join(item.employee_name if item.status == models.AssignmentStatus.WORK else f"{item.employee_name} · {_assignment_label(item)}" for item in assignments)
            lines.append(f"{_label(group)} ({len(assignments)}명): {names}")
        cell = ws.cell(row=row, column=column, value="\n".join(lines))
        cell.fill = _fill(next(iter(day.groups.keys()), "OTHER"))
        cell.font = Font(size=9)
        column += 1
        if column > 7:
            column = 1
            row += 1
    for row_index in range(5, row + 1):
        ws.row_dimensions[row_index].height = 115
    ws.print_area = f"A1:G{row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _apply_table_style(ws, row, 7)


def _build_employee_sheet(ws, calendar: schemas.MonthlyScheduleRead) -> None:
    _header(ws, "직원별 월간 근무현황")
    dates = [day.local_date for day in calendar.days]
    headers = ["직원", "근무일", "휴무일", "야간", *[item.isoformat() for item in dates]]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=2, column=column, value=header)
    employees = {item.employee_id: item.employee_name for day in calendar.days for item in day.assignments}
    by_employee = defaultdict(dict)
    for day in calendar.days:
        for item in day.assignments:
            by_employee[item.employee_id][day.local_date] = item
    for row_index, employee_id in enumerate(sorted(employees, key=lambda value: employees[value]), start=3):
        assignments = by_employee[employee_id]
        values = list(assignments.values())
        work_count = sum(1 for item in values if item.status == models.AssignmentStatus.WORK)
        off_count = len(values) - work_count
        night_count = sum(1 for item in values if item.category in {"NIGHT", "EARLY_NIGHT"})
        ws.cell(row=row_index, column=1, value=employees[employee_id])
        ws.cell(row=row_index, column=2, value=work_count)
        ws.cell(row=row_index, column=3, value=off_count)
        ws.cell(row=row_index, column=4, value=night_count)
        for column_index, local_date in enumerate(dates, start=5):
            item = assignments[local_date]
            cell = ws.cell(row=row_index, column=column_index, value=_assignment_label(item))
            cell.fill = _fill(item.category if item.status == models.AssignmentStatus.WORK else item.leave_type_code or item.status)
    ws.column_dimensions["A"].width = 18
    for column in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 11
    _apply_table_style(ws, len(employees) + 2, len(headers))


def _build_shift_count_sheet(ws, calendar: schemas.MonthlyScheduleRead) -> None:
    _header(ws, "근무유형별 배정 횟수")
    headers = ["근무유형", "분류", "배정 횟수"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=2, column=column, value=header)
    counts: Counter[tuple[str, str]] = Counter()
    for day in calendar.days:
        for item in day.assignments:
            key = (_assignment_label(item), _label(item.category if item.status == models.AssignmentStatus.WORK else item.leave_type_code or item.status))
            counts[key] += 1
    for row_index, ((name, category), count) in enumerate(sorted(counts.items()), start=3):
        ws.cell(row=row_index, column=1, value=name)
        ws.cell(row=row_index, column=2, value=category)
        ws.cell(row=row_index, column=3, value=count)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    _apply_table_style(ws, max(len(counts) + 2, 2), 3)


def _build_fairness_sheet(ws, calendar: schemas.MonthlyScheduleRead) -> None:
    _header(ws, "야간근무 및 휴무 공정성 통계")
    headers = ["직원", "총 근무", "야간", "조기", "주말근무", "휴무"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=2, column=column, value=header)
    employees = {item.employee_id: item.employee_name for day in calendar.days for item in day.assignments}
    by_employee = defaultdict(list)
    for day in calendar.days:
        local_date = date.fromisoformat(day.local_date.isoformat())
        for item in day.assignments:
            by_employee[item.employee_id].append((local_date, item))
    for row_index, employee_id in enumerate(sorted(employees, key=lambda value: employees[value]), start=3):
        values = by_employee[employee_id]
        work = sum(1 for _, item in values if item.status == models.AssignmentStatus.WORK)
        night = sum(1 for _, item in values if item.category in {"NIGHT", "EARLY_NIGHT"})
        early = sum(1 for _, item in values if item.category in {"EARLY_DAY", "EARLY_NIGHT"})
        weekend = sum(1 for local_date, item in values if local_date.weekday() >= 5 and item.status == models.AssignmentStatus.WORK)
        off = len(values) - work
        for column, value in enumerate([employees[employee_id], work, night, early, weekend, off], start=1):
            ws.cell(row=row_index, column=column, value=value)
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 14
    _apply_table_style(ws, len(employees) + 2, len(headers))


def _build_validation_sheet(ws, validation: schemas.ScheduleValidationRead) -> None:
    _header(ws, "검증 결과")
    ws["A2"] = "상태"
    ws["B2"] = "통과" if validation.ok else "확인 필요"
    headers = ["심각도", "코드", "날짜", "직원 ID", "근무유형 ID", "내용"]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=4, column=column, value=header)
    for row_index, issue in enumerate(validation.issues, start=5):
        values = [issue.severity, issue.code, issue.local_date.isoformat() if issue.local_date else "", issue.employee_id or "", issue.shift_type_id or "", issue.message]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=value)
    for column in range(1, 7):
        ws.column_dimensions[get_column_letter(column)].width = 20 if column != 6 else 48
    _apply_table_style(ws, max(len(validation.issues) + 4, 4), 6)
