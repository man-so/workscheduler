from collections.abc import Generator
from io import BytesIO
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app


engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db() -> Generator[Session, None, None]:
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


def setup_module() -> None:
    Base.metadata.create_all(bind=engine)
    app.dependency_overrides[get_db] = override_get_db


def teardown_module() -> None:
    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=engine)


client = TestClient(app)


def test_health() -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_phase_one_core_data_flow() -> None:
    team = client.post("/api/teams", json={"name": "검사실 A"}).json()

    employee_response = client.post(
        f"/api/teams/{team['id']}/employees",
        json={"employee_no": "EMP-001", "display_name": "직원 1"},
    )
    assert employee_response.status_code == 201
    employee = employee_response.json()

    shift_response = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={
            "code": "DAY",
            "name": "주간",
            "category": "DAY",
            "start_time": "09:00:00",
            "end_time": "18:00:00",
            "break_minutes": 60,
            "paid_minutes": 480,
        },
    )
    assert shift_response.status_code == 201
    shift = shift_response.json()

    rule_response = client.post(
        f"/api/teams/{team['id']}/work-rules",
        json={"name": "기본 규칙", "min_rest_minutes": 660, "max_consecutive_work_days": 5},
    )
    assert rule_response.status_code == 201

    leave_response = client.post(
        f"/api/teams/{team['id']}/leave-requests",
        json={
            "employee_id": employee["id"],
            "start_date": "2026-11-10",
            "end_date": "2026-11-10",
            "status": "APPROVED",
        },
    )
    assert leave_response.status_code == 201

    assignment_response = client.post(
        f"/api/teams/{team['id']}/assignments",
        json={
            "employee_id": employee["id"],
            "local_date": "2026-11-11",
            "status": "WORK",
            "shift_type_id": shift["id"],
        },
    )
    assert assignment_response.status_code == 201

    employees = client.get(f"/api/teams/{team['id']}/employees").json()
    assignments = client.get(f"/api/teams/{team['id']}/assignments").json()
    assert len(employees) == 1
    assert len(assignments) == 1


def test_setup_chat_preview_and_approval_with_overnight_shift() -> None:
    team = client.post("/api/teams", json={"name": "챗봇 테스트"}).json()
    session = client.get(f"/api/teams/{team['id']}/setup-sessions/active").json()

    response = client.post(
        f"/api/teams/{team['id']}/setup-sessions/{session['id']}/messages",
        json={
            "message": "직원은 김하나, 이둘. 주간 09:00-18:00, 야간 22:00-06:00. 최소 휴식 11시간, 연속근무 5일, 야간 후 휴무."
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["needs_approval"] is True
    assert body["pending_patch"]["shift_types"][1]["ends_next_day"] is True

    approval = client.post(f"/api/teams/{team['id']}/setup-sessions/{session['id']}/approve")
    assert approval.status_code == 200
    summary = approval.json()["summary"]
    assert len(summary["employees"]) == 2
    assert len(summary["shift_types"]) == 2
    assert summary["work_rules"][0]["min_rest_minutes"] == 660


def test_mixed_weekly_contracts_and_coverage() -> None:
    team = client.post("/api/teams", json={"name": "계약 테스트"}).json()
    emp4 = client.post(
        f"/api/teams/{team['id']}/employees",
        json={"employee_no": "EMP-4", "display_name": "주사일"},
    ).json()
    emp5 = client.post(
        f"/api/teams/{team['id']}/employees",
        json={"employee_no": "EMP-5", "display_name": "주오일"},
    ).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "DAY-MIX", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "18:00:00"},
    ).json()

    assert client.post(
        f"/api/teams/{team['id']}/employee-contracts",
        json={"employee_id": emp4["id"], "weekly_work_days": 4, "allowed_shift_type_ids_json": f"[{shift['id']}]"},
    ).status_code == 201
    assert client.post(
        f"/api/teams/{team['id']}/employee-contracts",
        json={"employee_id": emp5["id"], "weekly_work_days": 5, "allowed_shift_type_ids_json": f"[{shift['id']}]"},
    ).status_code == 201

    coverage = client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={
            "name": "평일 주간",
            "shift_type_id": shift["id"],
            "days_of_week_json": "[\"MON\",\"TUE\",\"WED\",\"THU\",\"FRI\"]",
            "min_count": 1,
            "target_count": 1,
            "max_count": 1,
        },
    )
    assert coverage.status_code == 201
    summary = client.get(f"/api/teams/{team['id']}/setup-summary").json()
    assert summary["conflicts"] == []


def test_conflicting_coverage_validation_and_bad_llm_json() -> None:
    from app.llm import validate_llm_json

    team = client.post("/api/teams", json={"name": "검증 테스트"}).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "DAY-CONFLICT", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "18:00:00"},
    ).json()
    response = client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={
            "name": "잘못된 필요 인원",
            "shift_type_id": shift["id"],
            "days_of_week_json": "[\"MON\"]",
            "min_count": 2,
            "target_count": 3,
            "max_count": 1,
        },
    )
    assert response.status_code == 422

    try:
        validate_llm_json("{\"intent\":\"setup\"}")
    except ValueError as exc:
        assert "missing fields" in str(exc)
    else:
        raise AssertionError("invalid LLM JSON should fail schema validation")


def test_generation_pipeline_persists_monthly_calendar(monkeypatch) -> None:
    from app import generation
    from app import main
    from app.agent_runner import AgentResult

    monkeypatch.setattr(generation, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(main, "submit_generation", lambda run_id, team_id, year, month, agent_id, timeout_seconds: generation.run_generation(run_id, team_id, year, month, agent_id, timeout_seconds))

    def fake_agent(agent_id, constraint_input, timeout_seconds, cancel_event=None):
        return AgentResult(payload=constraint_input, stdout="{}", stderr="", command=[agent_id])

    monkeypatch.setattr(generation, "run_agent", fake_agent)
    team = client.post("/api/teams", json={"name": "생성 파이프라인"}).json()
    employee = client.post(
        f"/api/teams/{team['id']}/employees",
        json={"employee_no": "EMP-GEN", "display_name": "가상 직원"},
    ).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "GEN-DAY", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "17:00:00"},
    ).json()
    assert client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={"name": "평일 주간", "shift_type_id": shift["id"], "days_of_week_json": "[\"MON\",\"TUE\",\"WED\",\"THU\",\"FRI\"]", "min_count": 1, "target_count": 1, "max_count": 1},
    ).status_code == 201

    started = client.post(f"/api/teams/{team['id']}/schedules/generate", json={"year": 2026, "month": 11, "agent_id": "codex", "timeout_seconds": 5})
    assert started.status_code == 202
    run_id = started.json()["id"]
    for _ in range(50):
        current = client.get(f"/api/teams/{team['id']}/generation-runs/{run_id}").json()
        if current["status"] in {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}:
            break
        time.sleep(0.02)
    assert current["status"] == "SUCCEEDED", current
    calendar = client.get(f"/api/teams/{team['id']}/schedules/2026/11")
    assert calendar.status_code == 200
    days = calendar.json()["days"]
    assert len(days) == 30
    assert all(len(day["assignments"]) == 1 for day in days)


def test_manual_calendar_edit_and_validation(monkeypatch) -> None:
    from app import generation
    from app import main
    from app.agent_runner import AgentResult

    monkeypatch.setattr(generation, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(main, "submit_generation", lambda run_id, team_id, year, month, agent_id, timeout_seconds: generation.run_generation(run_id, team_id, year, month, agent_id, timeout_seconds))

    def fake_agent(agent_id, constraint_input, timeout_seconds, cancel_event=None):
        return AgentResult(payload=constraint_input, stdout="{}", stderr="", command=[agent_id])

    monkeypatch.setattr(generation, "run_agent", fake_agent)
    team = client.post("/api/teams", json={"name": "수동수정 테스트"}).json()
    employee = client.post(
        f"/api/teams/{team['id']}/employees",
        json={"employee_no": "EMP-EDIT", "display_name": "수정 직원"},
    ).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "EDIT-DAY", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "17:00:00"},
    ).json()
    assert client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={"name": "평일 주간", "shift_type_id": shift["id"], "days_of_week_json": "[\"MON\"]", "min_count": 1, "target_count": 1, "max_count": 1},
    ).status_code == 201

    started = client.post(f"/api/teams/{team['id']}/schedules/generate", json={"year": 2026, "month": 11, "agent_id": "codex", "timeout_seconds": 5})
    run_id = started.json()["id"]
    for _ in range(50):
        current = client.get(f"/api/teams/{team['id']}/generation-runs/{run_id}").json()
        if current["status"] in {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}:
            break
        time.sleep(0.02)
    assert current["status"] == "SUCCEEDED", current

    update = client.patch(
        f"/api/teams/{team['id']}/schedules/2026/11/assignments",
        json={"employee_id": employee["id"], "local_date": "2026-11-02", "status": "OFF", "shift_type_id": None},
    )
    assert update.status_code == 200
    edited_day = next(day for day in update.json()["days"] if day["local_date"] == "2026-11-02")
    assert edited_day["assignments"][0]["status"] == "OFF"

    validation = client.get(f"/api/teams/{team['id']}/schedules/2026/11/validation")
    assert validation.status_code == 200
    body = validation.json()
    assert body["ok"] is False
    assert any(issue["code"] == "COVERAGE_BELOW_MIN" for issue in body["issues"])


def test_confirm_blocks_direct_edit_and_clone_creates_new_draft(monkeypatch) -> None:
    from app import generation
    from app import main
    from app.agent_runner import AgentResult

    monkeypatch.setattr(generation, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(main, "submit_generation", lambda run_id, team_id, year, month, agent_id, timeout_seconds: generation.run_generation(run_id, team_id, year, month, agent_id, timeout_seconds))
    monkeypatch.setattr(generation, "run_agent", lambda agent_id, constraint_input, timeout_seconds, cancel_event=None: AgentResult(payload=constraint_input, stdout="{}", stderr="", command=[agent_id]))

    team = client.post("/api/teams", json={"name": "확정 테스트"}).json()
    employee = client.post(f"/api/teams/{team['id']}/employees", json={"employee_no": "EMP-CONFIRM", "display_name": "확정 직원"}).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "CONFIRM-DAY", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "17:00:00"},
    ).json()
    assert client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={"name": "월요일 주간", "shift_type_id": shift["id"], "days_of_week_json": "[\"MON\"]", "min_count": 1, "target_count": 1, "max_count": 1},
    ).status_code == 201

    started = client.post(f"/api/teams/{team['id']}/schedules/generate", json={"year": 2026, "month": 11, "agent_id": "codex", "timeout_seconds": 5})
    run_id = started.json()["id"]
    for _ in range(50):
        current = client.get(f"/api/teams/{team['id']}/generation-runs/{run_id}").json()
        if current["status"] in {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}:
            break
        time.sleep(0.02)
    assert current["status"] == "SUCCEEDED", current

    confirmed = client.post(f"/api/teams/{team['id']}/schedules/2026/11/confirm", json={"approve_soft_issues": True})
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "CONFIRMED"

    blocked = client.patch(
        f"/api/teams/{team['id']}/schedules/2026/11/assignments",
        json={"employee_id": employee["id"], "local_date": "2026-11-02", "status": "OFF", "shift_type_id": None},
    )
    assert blocked.status_code == 409

    cloned = client.post(f"/api/teams/{team['id']}/schedules/2026/11/versions/{confirmed.json()['id']}/clone")
    assert cloned.status_code == 201
    assert cloned.json()["status"] == "DRAFT"
    assert cloned.json()["is_active"] is True

    edited = client.patch(
        f"/api/teams/{team['id']}/schedules/2026/11/assignments",
        json={"employee_id": employee["id"], "local_date": "2026-11-02", "status": "OFF", "shift_type_id": None},
    )
    assert edited.status_code == 200


def test_confirm_blocks_hard_constraint_errors(monkeypatch) -> None:
    from app import generation
    from app import main
    from app.agent_runner import AgentResult

    monkeypatch.setattr(generation, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(main, "submit_generation", lambda run_id, team_id, year, month, agent_id, timeout_seconds: generation.run_generation(run_id, team_id, year, month, agent_id, timeout_seconds))
    monkeypatch.setattr(generation, "run_agent", lambda agent_id, constraint_input, timeout_seconds, cancel_event=None: AgentResult(payload=constraint_input, stdout="{}", stderr="", command=[agent_id]))

    team = client.post("/api/teams", json={"name": "Hard 차단 테스트"}).json()
    employee = client.post(f"/api/teams/{team['id']}/employees", json={"employee_no": "EMP-HARD", "display_name": "Hard 직원"}).json()
    shift = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "HARD-DAY", "name": "주간", "category": "DAY", "start_time": "09:00:00", "end_time": "17:00:00"},
    ).json()
    client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={"name": "월요일 주간", "shift_type_id": shift["id"], "days_of_week_json": "[\"MON\"]", "min_count": 1, "target_count": 1, "max_count": 1},
    )
    started = client.post(f"/api/teams/{team['id']}/schedules/generate", json={"year": 2026, "month": 11, "agent_id": "codex", "timeout_seconds": 5})
    run_id = started.json()["id"]
    for _ in range(50):
        current = client.get(f"/api/teams/{team['id']}/generation-runs/{run_id}").json()
        if current["status"] in {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}:
            break
        time.sleep(0.02)
    assert current["status"] == "SUCCEEDED", current
    client.patch(
        f"/api/teams/{team['id']}/schedules/2026/11/assignments",
        json={"employee_id": employee["id"], "local_date": "2026-11-02", "status": "OFF", "shift_type_id": None},
    )
    blocked = client.post(f"/api/teams/{team['id']}/schedules/2026/11/confirm", json={"approve_soft_issues": True})
    assert blocked.status_code == 422
    assert blocked.json()["detail"][0]["code"] == "COVERAGE_BELOW_MIN"


def test_excel_export_contains_calendar_summaries_and_all_employees(monkeypatch) -> None:
    from app import generation
    from app import main
    from app.agent_runner import AgentResult

    monkeypatch.setattr(generation, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(main, "submit_generation", lambda run_id, team_id, year, month, agent_id, timeout_seconds: generation.run_generation(run_id, team_id, year, month, agent_id, timeout_seconds))
    monkeypatch.setattr(generation, "run_agent", lambda agent_id, constraint_input, timeout_seconds, cancel_event=None: AgentResult(payload=constraint_input, stdout="{}", stderr="", command=[agent_id]))

    team = client.post("/api/teams", json={"name": "Excel 테스트"}).json()
    employees = [
        client.post(f"/api/teams/{team['id']}/employees", json={"employee_no": "EMP-X1", "display_name": "엑셀 하나"}).json(),
        client.post(f"/api/teams/{team['id']}/employees", json={"employee_no": "EMP-X2", "display_name": "엑셀 둘"}).json(),
    ]
    special = client.post(
        f"/api/teams/{team['id']}/shift-types",
        json={"code": "SPECIAL", "name": "특수근무", "category": "SPECIAL", "start_time": "13:00:00", "end_time": "21:00:00", "color": "#FDE68A"},
    ).json()
    client.post(
        f"/api/teams/{team['id']}/coverage-requirements",
        json={"name": "월초월말 특수", "shift_type_id": special["id"], "days_of_week_json": "[\"MON\",\"THU\"]", "min_count": 1, "target_count": 1, "max_count": 1},
    )
    started = client.post(f"/api/teams/{team['id']}/schedules/generate", json={"year": 2026, "month": 12, "agent_id": "codex", "timeout_seconds": 5})
    run_id = started.json()["id"]
    for _ in range(50):
        current = client.get(f"/api/teams/{team['id']}/generation-runs/{run_id}").json()
        if current["status"] in {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}:
            break
        time.sleep(0.02)
    assert current["status"] == "SUCCEEDED", current

    response = client.get(f"/api/teams/{team['id']}/schedules/2026/12/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["월간 달력", "직원별 현황", "근무유형 통계", "공정성 통계", "검증 결과"]
    calendar_text = "\n".join(str(cell.value or "") for row in workbook["월간 달력"].iter_rows() for cell in row)
    employee_text = "\n".join(str(cell.value or "") for row in workbook["직원별 현황"].iter_rows() for cell in row)
    stats_text = "\n".join(str(cell.value or "") for row in workbook["근무유형 통계"].iter_rows() for cell in row)
    for employee in employees:
        assert employee["display_name"] in calendar_text
        assert employee["display_name"] in employee_text
    assert "31일" in calendar_text
    assert "특수근무" in stats_text
