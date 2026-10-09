from __future__ import annotations

import argparse
import json
import os
import sqlite3
import sys
import threading
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

os.environ.setdefault("DATABASE_URL", f"sqlite:///{(ROOT / 'artifacts' / 'phase6' / 'phase6.db').as_posix()}")
os.environ.setdefault("BACKEND_CORS_ORIGINS", "http://localhost:3000")

from app import models  # noqa: E402
from app.constraints import build_constraint_payload  # noqa: E402
from app.database import Base  # noqa: E402
from app.excel_export import build_schedule_workbook  # noqa: E402
from app.generation import _create_schedule  # noqa: E402
from app.main import build_monthly_schedule_read, validate_schedule_version  # noqa: E402
from app.setup_logic import dumps  # noqa: E402
from app.solver import solve_month  # noqa: E402


@dataclass
class ScenarioResult:
    code: str
    name: str
    status: str
    solver_status: str | None
    elapsed_seconds: float
    validation_ok: bool | None
    hard_errors: int
    soft_warnings: int
    fairness: dict[str, Any]
    excel: dict[str, Any] | None
    notes: list[str]


def fresh_engine(db_path: Path):
    if db_path.exists():
        db_path.unlink()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    engine = create_engine(f"sqlite:///{db_path.as_posix()}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    return engine


def create_team(db, name: str) -> models.Team:
    team = models.Team(name=name, timezone="Asia/Seoul")
    db.add(team)
    db.flush()
    return team


def create_employees(db, team: models.Team, count: int, prefix: str) -> list[models.Employee]:
    employees = []
    for index in range(1, count + 1):
        employee = models.Employee(team_id=team.id, employee_no=f"{prefix}-{index:02d}", display_name=f"{prefix} 직원 {index:02d}")
        db.add(employee)
        employees.append(employee)
    db.flush()
    return employees


def create_shift(db, team: models.Team, code: str, name: str, category: str, start: str, end: str, color: str | None = None) -> models.ShiftType:
    shift = models.ShiftType(
        team_id=team.id,
        code=code,
        name=name,
        category=category,
        start_time=datetime.strptime(start, "%H:%M").time(),
        end_time=datetime.strptime(end, "%H:%M").time(),
        ends_next_day=end <= start,
        paid_minutes=480,
        color=color,
        is_work=True,
        is_active=True,
    )
    db.add(shift)
    db.flush()
    return shift


def add_coverage(db, team: models.Team, shift: models.ShiftType, days: list[str], count: int, name: str | None = None) -> None:
    db.add(models.CoverageRequirement(
        team_id=team.id,
        name=name or f"{shift.name} 필요 인원",
        shift_type_id=shift.id,
        days_of_week_json=dumps(days),
        min_count=count,
        target_count=count,
        max_count=count,
        priority=100,
        is_active=True,
    ))


def add_contract(db, team: models.Team, employee: models.Employee, weekly_days: int, shift_ids: list[int]) -> None:
    db.add(models.EmployeeContract(
        team_id=team.id,
        employee_id=employee.id,
        weekly_work_days=weekly_days,
        allowed_shift_type_ids_json=dumps(shift_ids),
    ))


def add_rule(db, team: models.Team, name: str = "기본 규칙") -> None:
    db.add(models.WorkRule(team_id=team.id, name=name, min_rest_minutes=660, max_consecutive_work_days=6, max_consecutive_night_shifts=3, night_requires_next_day_off=True))


def build_scenario(db, code: str) -> tuple[models.Team, int, int]:
    year, month = 2026, 11
    if code == "A":
        team = create_team(db, "A. 8명 주간 야간 휴무")
        employees = create_employees(db, team, 8, "A")
        day = create_shift(db, team, "DAY", "주간", "DAY", "09:00", "17:00")
        night = create_shift(db, team, "NIGHT", "야간", "NIGHT", "22:00", "06:00")
        add_coverage(db, team, day, ["MON", "TUE", "WED", "THU", "FRI"], 3)
        add_coverage(db, team, night, ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"], 1)
        for employee in employees:
            add_contract(db, team, employee, 5, [day.id, night.id])
        add_rule(db, team)
    elif code == "B":
        team = create_team(db, "B. 8명 주4 주5 혼합")
        employees = create_employees(db, team, 8, "B")
        day = create_shift(db, team, "DAY", "주간", "DAY", "09:00", "17:00")
        add_coverage(db, team, day, ["MON", "TUE", "WED", "THU", "FRI"], 4)
        for index, employee in enumerate(employees):
            add_contract(db, team, employee, 4 if index < 4 else 5, [day.id])
        add_rule(db, team)
    elif code == "C":
        team = create_team(db, "C. 15명 조기 야간조기")
        employees = create_employees(db, team, 15, "C")
        day = create_shift(db, team, "DAY", "주간", "DAY", "09:00", "17:00")
        early = create_shift(db, team, "EARLY", "조기", "EARLY_DAY", "06:00", "14:00")
        night = create_shift(db, team, "NIGHT", "야간", "NIGHT", "22:00", "06:00")
        early_night = create_shift(db, team, "EARLY_NIGHT", "야간 조기", "EARLY_NIGHT", "20:00", "04:00")
        add_coverage(db, team, day, ["MON", "TUE", "WED", "THU", "FRI"], 4)
        add_coverage(db, team, early, ["MON", "TUE", "WED", "THU", "FRI"], 2)
        add_coverage(db, team, night, ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"], 1)
        add_coverage(db, team, early_night, ["MON", "WED", "FRI"], 1)
        for employee in employees:
            add_contract(db, team, employee, 5, [day.id, early.id, night.id, early_night.id])
        add_rule(db, team)
    elif code == "D":
        team = create_team(db, "D. 25명 근무유형 10개")
        employees = create_employees(db, team, 25, "D")
        shifts = []
        categories = ["DAY", "EARLY_DAY", "DAY", "OTHER", "NIGHT", "EARLY_NIGHT", "OTHER", "DAY", "OTHER", "NIGHT"]
        times = [("08:00", "16:00"), ("06:00", "14:00"), ("09:00", "17:00"), ("10:00", "18:00"), ("22:00", "06:00"), ("20:00", "04:00"), ("12:00", "20:00"), ("07:00", "15:00"), ("13:00", "21:00"), ("23:00", "07:00")]
        for index in range(10):
            shifts.append(create_shift(db, team, f"S{index+1}", f"근무유형 {index+1}", categories[index], *times[index]))
        for shift in shifts[:8]:
            add_coverage(db, team, shift, ["MON", "TUE", "WED", "THU", "FRI"], 1)
        for shift in shifts[8:]:
            add_coverage(db, team, shift, ["SAT", "SUN"], 1)
        for employee in employees:
            add_contract(db, team, employee, 5, [shift.id for shift in shifts])
        add_rule(db, team)
    elif code == "E":
        team = create_team(db, "E. 휴무 신청 집중")
        employees = create_employees(db, team, 10, "E")
        day = create_shift(db, team, "DAY", "주간", "DAY", "09:00", "17:00")
        add_coverage(db, team, day, ["MON", "TUE", "WED", "THU", "FRI"], 3)
        for employee in employees:
            add_contract(db, team, employee, 5, [day.id])
        for employee in employees[:6]:
            db.add(models.Availability(team_id=team.id, employee_id=employee.id, local_date=date(year, month, 16), availability_type="UNAVAILABLE", source="MANUAL"))
            db.add(models.Availability(team_id=team.id, employee_id=employee.id, local_date=date(year, month, 17), availability_type="PREFERRED_OFF", source="MANUAL"))
        add_rule(db, team)
    elif code == "F":
        team = create_team(db, "F. Hard Constraint 충돌")
        employees = create_employees(db, team, 3, "F")
        day = create_shift(db, team, "DAY", "주간", "DAY", "09:00", "17:00")
        add_coverage(db, team, day, ["MON"], 4)
        for employee in employees:
            add_contract(db, team, employee, 5, [day.id])
        add_rule(db, team)
    else:
        raise ValueError(code)
    db.commit()
    return team, year, month


def fairness(calendar) -> dict[str, Any]:
    by_employee: dict[str, Counter[str]] = defaultdict(Counter)
    for day in calendar.days:
        for item in day.assignments:
            by_employee[item.employee_name]["work"] += 1 if item.status == "WORK" else 0
            by_employee[item.employee_name]["off"] += 1 if item.status != "WORK" else 0
            by_employee[item.employee_name]["night"] += 1 if item.category in {"NIGHT", "EARLY_NIGHT"} else 0
            by_employee[item.employee_name]["early"] += 1 if item.category in {"EARLY_DAY", "EARLY_NIGHT"} else 0
    night_counts = [counts["night"] for counts in by_employee.values()]
    work_counts = [counts["work"] for counts in by_employee.values()]
    return {
        "employees": len(by_employee),
        "night_min": min(night_counts) if night_counts else 0,
        "night_max": max(night_counts) if night_counts else 0,
        "work_min": min(work_counts) if work_counts else 0,
        "work_max": max(work_counts) if work_counts else 0,
    }


def validate_excel(blob: BytesIO, calendar) -> dict[str, Any]:
    wb = load_workbook(blob)
    expected = {"월간 달력", "직원별 현황", "근무유형 통계", "공정성 통계", "검증 결과"}
    calendar_text = "\n".join(str(cell.value or "") for row in wb["월간 달력"].iter_rows() for cell in row)
    employee_names = sorted({item.employee_name for day in calendar.days for item in day.assignments})
    missing = [name for name in employee_names if name not in calendar_text]
    days_missing = [day.local_date.isoformat() for day in calendar.days if f"{int(day.local_date.strftime('%d'))}일" not in calendar_text]
    return {
        "open_ok": True,
        "sheets_ok": set(wb.sheetnames) == expected,
        "all_employees_in_calendar": not missing,
        "missing_employees": missing,
        "all_days_present": not days_missing,
        "missing_days": days_missing,
        "print_area": wb["월간 달력"].print_area,
        "orientation": wb["월간 달력"].page_setup.orientation,
        "fit_to_width": wb["월간 달력"].page_setup.fitToWidth,
    }


def run_scenario(db, code: str, output_dir: Path) -> ScenarioResult:
    team, year, month = build_scenario(db, code)
    start = time.perf_counter()
    base = build_constraint_payload(db, team, year, month)
    result = solve_month(base, threading.Event(), 30)
    elapsed = time.perf_counter() - start
    notes: list[str] = []
    if result.status != "SUCCEEDED":
        return ScenarioResult(code, team.name, result.status, result.solver_status, elapsed, None, 0, 0, {}, None, result.conflicts)
    result_data = _create_schedule(db, team.id, year, month, result.solver_status, base, result.assignments)
    schedule = db.get(models.Schedule, result_data["schedule_id"])
    version = db.get(models.ScheduleVersion, result_data["version_id"])
    calendar = build_monthly_schedule_read(db, schedule, version)
    validation = validate_schedule_version(db, schedule, version)
    hard_errors = sum(1 for issue in validation.issues if issue.severity == "ERROR")
    soft_warnings = sum(1 for issue in validation.issues if issue.severity == "WARN")
    excel_blob = build_schedule_workbook(team, version, calendar, validation)
    excel_path = output_dir / f"scenario_{code}.xlsx"
    excel_path.write_bytes(excel_blob.getvalue())
    excel_result = validate_excel(BytesIO(excel_blob.getvalue()), calendar)
    return ScenarioResult(code, team.name, result.status, result.solver_status, elapsed, validation.ok, hard_errors, soft_warnings, fairness(calendar), excel_result, notes)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", default=str(ROOT / "artifacts" / "phase6" / "phase6.db"))
    parser.add_argument("--out", default=str(ROOT / "artifacts" / "phase6"))
    args = parser.parse_args()
    db_path = Path(args.db)
    output_dir = Path(args.out)
    output_dir.mkdir(parents=True, exist_ok=True)
    engine = fresh_engine(db_path)
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    results = []
    with SessionLocal() as db:
        for code in ["A", "B", "C", "D", "E", "F"]:
            results.append(run_scenario(db, code, output_dir).__dict__)
    report = {
        "database": str(db_path),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "scenarios": results,
    }
    report_path = output_dir / "phase6_scenarios.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
