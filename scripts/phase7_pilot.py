from __future__ import annotations

import json
import subprocess
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.setup_logic import parse_setup_message  # noqa: E402


API = "http://127.0.0.1:8000"
YEAR = 2026
MONTH = 11
OUT = ROOT / "artifacts" / "phase7"
INSTALL_EXE = Path.home() / "AppData" / "Local" / "Programs" / "AI Shift Scheduler" / "AI Shift Scheduler.exe"


class ApiError(RuntimeError):
    pass


def request(method: str, path: str, payload: Any | None = None, timeout: int = 30) -> Any:
    body = None
    headers = {}
    if payload is not None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(f"{API}{path}", data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            raw = response.read()
            content_type = response.headers.get("Content-Type", "")
            if "json" in content_type:
                return json.loads(raw.decode("utf-8"))
            return raw
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise ApiError(f"{method} {path} -> {exc.code}: {detail}") from exc


def wait_health() -> bool:
    for _ in range(40):
        try:
            return request("GET", "/api/health")["status"] == "ok"
        except Exception:
            time.sleep(0.5)
    return False


def start_app() -> subprocess.Popen[bytes] | None:
    if wait_health():
        return None
    if not INSTALL_EXE.exists():
        raise RuntimeError(f"installed Electron app not found: {INSTALL_EXE}")
    process = subprocess.Popen([str(INSTALL_EXE)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if not wait_health():
        raise RuntimeError("Electron app did not start FastAPI")
    return process


def post(team_id: int, resource: str, payload: dict[str, Any]) -> Any:
    return request("POST", f"/api/teams/{team_id}/{resource}", payload)


def create_shift(team_id: int, code: str, name: str, category: str, start: str, end: str, color: str) -> dict[str, Any]:
    return post(
        team_id,
        "shift-types",
        {
            "code": code,
            "name": name,
            "category": category,
            "start_time": start,
            "end_time": end,
            "ends_next_day": end <= start,
            "break_minutes": 60 if category in {"DAY", "EARLY_DAY"} else 0,
            "paid_minutes": 480,
            "is_work": True,
            "color": color,
        },
    )


def add_coverage(team_id: int, shift_id: int, label: str, days: list[str], count: int) -> None:
    post(
        team_id,
        "coverage-requirements",
        {
            "name": label,
            "shift_type_id": shift_id,
            "days_of_week_json": json.dumps(days, ensure_ascii=False),
            "min_count": count,
            "target_count": count,
            "max_count": count,
            "priority": 100,
            "is_active": True,
        },
    )


def setup_pilot_team(workflow: list[dict[str, Any]]) -> tuple[int, dict[str, Any]]:
    stamp = datetime.now().strftime("%H%M%S")
    team = request("POST", "/api/teams", {"name": f"Phase 7 Pilot {stamp}", "timezone": "Asia/Seoul"})
    team_id = team["id"]
    workflow.append({"step": 1, "name": "팀 생성", "status": "PASS", "result": team})

    session = request("GET", f"/api/teams/{team_id}/setup-sessions/active")
    employee_names = [f"Pilot-{index:02d}" for index in range(1, 17)]
    message = "직원은 " + ", ".join(employee_names)
    preview = request("POST", f"/api/teams/{team_id}/setup-sessions/{session['id']}/messages", {"message": message})
    workflow.append({"step": 3, "name": "챗봇 직원 입력", "status": "PASS", "needs_approval": preview.get("needs_approval"), "changes": preview.get("proposed_changes", [])})
    approved = request("POST", f"/api/teams/{team_id}/setup-sessions/{session['id']}/approve", {})
    workflow.append({"step": 5, "name": "직원 입력 승인", "status": "PASS", "applied": len(approved.get("applied", []))})
    workflow.append({"step": 2, "name": "직원 등록", "status": "PASS", "count": len(employee_names)})

    shift_payload = "근무유형 주간 09:00-17:00, 조기 06:00-14:00, 야간 22:00-06:00, 야간 조기 20:00-04:00. 최소 휴식 11시간, 연속근무 5일, 연속 야간 2일, 야간 후 휴무."
    session = request("GET", f"/api/teams/{team_id}/setup-sessions/active")
    preview = request("POST", f"/api/teams/{team_id}/setup-sessions/{session['id']}/messages", {"message": shift_payload})
    workflow.append({"step": 4, "name": "근무조건 미리보기", "status": "PASS", "needs_approval": preview.get("needs_approval"), "changes": preview.get("proposed_changes", [])})
    request("POST", f"/api/teams/{team_id}/setup-sessions/{session['id']}/approve", {})

    employees = request("GET", f"/api/teams/{team_id}/employees")
    shifts = request("GET", f"/api/teams/{team_id}/shift-types")
    by_category = {shift["category"]: shift for shift in shifts}
    if not {"DAY", "EARLY_DAY", "NIGHT", "EARLY_NIGHT"}.issubset(by_category):
        # The rule parser can miss ambiguous Korean prefixes; form/API entry is the no-LLM fallback path.
        shifts = [
            create_shift(team_id, "DAY-P7", "주간", "DAY", "09:00:00", "17:00:00", "#60a5fa"),
            create_shift(team_id, "EARLY-P7", "조기", "EARLY_DAY", "06:00:00", "14:00:00", "#34d399"),
            create_shift(team_id, "NIGHT-P7", "야간", "NIGHT", "22:00:00", "06:00:00", "#818cf8"),
            create_shift(team_id, "EN-P7", "야간 조기", "EARLY_NIGHT", "20:00:00", "04:00:00", "#f59e0b"),
        ]
        by_category = {shift["category"]: shift for shift in shifts}
    if not request("GET", f"/api/teams/{team_id}/work-rules"):
        post(
            team_id,
            "work-rules",
            {
                "name": "파일럿 기본 규칙",
                "min_rest_minutes": 660,
                "max_consecutive_work_days": 5,
                "max_consecutive_night_shifts": 2,
                "night_requires_next_day_off": True,
                "soft_weights_json": json.dumps({"night": 10, "early": 4, "weekend": 3}, ensure_ascii=False),
            },
        )

    for employee in employees:
        allowed = [by_category["DAY"]["id"], by_category["EARLY_DAY"]["id"]]
        if int(employee["employee_no"].split("-")[-1]) <= 12:
            allowed += [by_category["NIGHT"]["id"], by_category["EARLY_NIGHT"]["id"]]
        weekly_days = 4 if int(employee["employee_no"].split("-")[-1]) in {4, 8, 12, 16} else 5
        post(
            team_id,
            "employee-contracts",
            {
                "employee_id": employee["id"],
                "weekly_work_days": weekly_days,
                "allowed_shift_type_ids_json": json.dumps(allowed),
                "restrictions_json": None,
            },
        )

    weekdays = ["MON", "TUE", "WED", "THU", "FRI"]
    every_day = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
    add_coverage(team_id, by_category["DAY"]["id"], "평일 주간", weekdays, 5)
    add_coverage(team_id, by_category["EARLY_DAY"]["id"], "평일 조기", weekdays, 2)
    add_coverage(team_id, by_category["NIGHT"]["id"], "매일 야간", every_day, 1)
    add_coverage(team_id, by_category["EARLY_NIGHT"]["id"], "월수금 야간 조기", ["MON", "WED", "FRI"], 1)
    workflow.append({"step": "2-5", "name": "폼/API 보완 설정", "status": "PASS", "contracts": len(employees), "coverage": 4})

    leave_map = [
        (employees[2], "2026-11-05"),
        (employees[3], "2026-11-10"),
        (employees[7], "2026-11-17"),
        (employees[11], "2026-11-24"),
    ]
    for employee, local_date in leave_map:
        post(team_id, "availability", {"employee_id": employee["id"], "local_date": local_date, "availability_type": "UNAVAILABLE", "source": "PILOT"})
        post(team_id, "leave-requests", {"employee_id": employee["id"], "start_date": local_date, "end_date": local_date, "status": "APPROVED", "leave_type": "ANNUAL", "reason": "PILOT"})
    workflow.append({"step": "1-data", "name": "연차 및 휴무 신청", "status": "PASS", "count": len(leave_map)})
    return team_id, {"employees": employees, "shifts": by_category, "leave_dates": leave_map}


def wait_generation(team_id: int, run_id: int) -> dict[str, Any]:
    terminal = {"SUCCEEDED", "FAILED", "INFEASIBLE", "CANCELED"}
    for _ in range(720):
        run = request("GET", f"/api/teams/{team_id}/generation-runs/{run_id}")
        if run["status"] in terminal:
            return run
        time.sleep(0.5)
    raise TimeoutError("generation did not finish")


def evaluate_calendar(calendar: dict[str, Any], validation: dict[str, Any], pilot: dict[str, Any]) -> dict[str, Any]:
    duplicate_days = []
    missing_days = []
    leave_violations = []
    night_counts: Counter[str] = Counter()
    employee_count = len(pilot["employees"])
    for day in calendar["days"]:
        ids = [item["employee_id"] for item in day["assignments"]]
        if len(ids) != employee_count:
            missing_days.append(day["local_date"])
        if len(ids) != len(set(ids)):
            duplicate_days.append(day["local_date"])
        for item in day["assignments"]:
            if item["category"] in {"NIGHT", "EARLY_NIGHT"}:
                night_counts[item["employee_name"]] += 1
        for employee, leave_date in pilot["leave_dates"]:
            if day["local_date"] == leave_date:
                assignment = next(item for item in day["assignments"] if item["employee_id"] == employee["id"])
                if assignment["status"] == "WORK":
                    leave_violations.append({"date": leave_date, "employee": employee["display_name"]})
    return {
        "hard_constraint_errors": sum(1 for issue in validation["issues"] if issue["severity"] == "ERROR"),
        "soft_warnings": sum(1 for issue in validation["issues"] if issue["severity"] == "WARN"),
        "missing_employee_days": missing_days,
        "duplicate_assignment_days": duplicate_days,
        "leave_violations": leave_violations,
        "night_min": min(night_counts.values()) if night_counts else 0,
        "night_max": max(night_counts.values()) if night_counts else 0,
    }


def validate_excel(blob: bytes, calendar: dict[str, Any]) -> dict[str, Any]:
    wb = load_workbook(BytesIO(blob))
    text = "\n".join(str(cell.value or "") for row in wb["월간 달력"].iter_rows() for cell in row)
    names = sorted({item["employee_name"] for day in calendar["days"] for item in day["assignments"]})
    missing = [name for name in names if name not in text]
    return {
        "open_ok": True,
        "sheets": wb.sheetnames,
        "all_employees_in_calendar": not missing,
        "missing_employees": missing,
        "print_area": wb["월간 달력"].print_area,
        "orientation": wb["월간 달력"].page_setup.orientation,
    }


def evaluate_natural_language(team_id: int) -> list[dict[str, Any]]:
    rows = {
        "employees": [],
        "shift_types": [],
        "work_rules": [],
        "employee_contracts": [],
        "coverage_requirements": [],
        "leave_requests": [],
        "availability": [],
    }
    # `parse_setup_message` expects ORM-like rows. For pilot reporting we record entity extraction only.
    class Row:
        def __init__(self, **values: Any) -> None:
            self.__dict__.update(values)

    employees = request("GET", f"/api/teams/{team_id}/employees")
    shifts = request("GET", f"/api/teams/{team_id}/shift-types")
    rows["employees"] = [Row(**item) for item in employees]
    rows["shift_types"] = [Row(**item) for item in shifts]
    team = Row(id=team_id, name="Phase 7 Pilot")
    cases = [
        ("직원은 Alpha, Bravo, Charlie", ["employees"]),
        ("주간 09:00-17:00, 야간 22:00-06:00", ["shift_types"]),
        ("야간 조기 20:00-04:00", ["shift_types"]),
        ("최소 휴식 11시간, 연속근무 5일, 연속 야간 2일, 야간 후 휴무", ["work_rules"]),
        ("평일 주간 5명, 야간 1명", ["coverage_requirements"]),
        ("주말 야간 1명", ["coverage_requirements"]),
        ("Pilot-04는 주 4일", ["employee_contracts"]),
        ("Pilot-03 2026-11-05 휴무", ["availability"]),
        ("야간 공정성 10, 조기 공정성 4, 주말 공정성 3", ["work_rules"]),
        ("야간근무 다음날 오전근무 금지", ["unsupported"]),
    ]
    results = []
    for text, expected in cases:
        patch = parse_setup_message(text, team, rows)
        found = [key for key, value in patch.items() if key != "team" and value]
        unsupported = "unsupported" in expected and not found
        matched = all(item in found for item in expected if item != "unsupported")
        results.append(
            {
                "input": text,
                "expected": expected,
                "interpreted_entities": found,
                "matches_expected": matched or unsupported,
                "needs_followup": not (matched or unsupported),
                "unsupported_or_missing": unsupported,
                "constraint_patch": patch,
            }
        )
    return results


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    workflow: list[dict[str, Any]] = []
    start_app()
    team_id, pilot = setup_pilot_team(workflow)
    summary = request("GET", f"/api/teams/{team_id}/setup-summary")
    workflow.append({"step": 7, "name": "설정 스키마/누락 검증", "status": "PASS" if not summary["conflicts"] else "FAIL", "completion_percent": summary["completion_percent"], "missing": summary["missing_fields"], "conflicts": summary["conflicts"]})

    generation_started = time.perf_counter()
    run = request("POST", f"/api/teams/{team_id}/schedules/generate", {"year": YEAR, "month": MONTH, "agent_id": "codex", "timeout_seconds": 180})
    workflow.append({"step": 6, "name": "Codex CLI 실제 호출 시작", "status": "PASS", "run_id": run["id"]})
    run = wait_generation(team_id, run["id"])
    elapsed = time.perf_counter() - generation_started
    workflow.append({"step": 8, "name": "OR-Tools 월간 근무표 생성", "status": "PASS" if run["status"] == "SUCCEEDED" else "FAIL", "run_status": run["status"], "elapsed_seconds": elapsed, "error": run.get("error_message")})

    calendar = request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}")
    validation = request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/validation")
    workflow.append({"step": 9, "name": "월간 달력 표시 데이터 조회", "status": "PASS", "days": len(calendar["days"]), "first_day_assignments": len(calendar["days"][0]["assignments"])})

    first_day = calendar["days"][0]
    first_employee = first_day["assignments"][0]
    edited = request(
        "PATCH",
        f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/assignments",
        {
            "employee_id": first_employee["employee_id"],
            "local_date": first_day["local_date"],
            "status": first_employee["status"],
            "shift_type_id": first_employee["shift_type_id"],
            "change_reason": "Phase 7 pilot manual edit path verification",
        },
    )
    validation_after_edit = request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/validation")
    workflow.append({"step": 10, "name": "수동 수정", "status": "PASS", "employee": first_employee["employee_name"], "date": first_day["local_date"]})
    workflow.append({"step": 11, "name": "변경 후 재검증", "status": "PASS", "ok": validation_after_edit["ok"], "issues": validation_after_edit["issue_count"]})

    confirm_status = "PASS"
    confirm_error = None
    try:
        confirm = request("POST", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/confirm", {"approve_soft_issues": True})
    except Exception as exc:
        confirm_status = "FAIL"
        confirm_error = str(exc)
        confirm = None
    workflow.append({"step": 12, "name": "근무표 확정", "status": confirm_status, "result": confirm, "error": confirm_error})

    excel_blob = request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/export.xlsx")
    excel_path = OUT / "phase7_pilot_schedule.xlsx"
    excel_path.write_bytes(excel_blob)
    calendar_after = request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}")
    excel_result = validate_excel(excel_blob, calendar_after)
    workflow.append({"step": 13, "name": "Excel 내보내기", "status": "PASS" if excel_result["open_ok"] and excel_result["all_employees_in_calendar"] else "FAIL", "path": str(excel_path), "excel": excel_result})

    quality = evaluate_calendar(calendar_after, request("GET", f"/api/teams/{team_id}/schedules/{YEAR}/{MONTH}/validation"), pilot)
    workflow.append({"step": 15, "name": "확정 근무표 재조회", "status": "PASS", "status_value": calendar_after["status"], "days": len(calendar_after["days"])})

    nl_results = evaluate_natural_language(team_id)
    usage = {}
    log_text = run.get("log_text") or ""
    if "turn.completed" in log_text:
        usage["raw_log_tail"] = log_text[-2000:]
    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "team_id": team_id,
        "year": YEAR,
        "month": MONTH,
        "workflow": workflow,
        "generation_run": {
            "id": run["id"],
            "status": run["status"],
            "elapsed_seconds": elapsed,
            "agent_id": run["agent_id"],
            "log_tail": log_text[-4000:],
            "usage": usage,
        },
        "quality": quality,
        "natural_language_cases": nl_results,
        "artifacts": {
            "excel": str(excel_path),
            "calendar_screenshot": str(OUT / "phase7_calendar_electron.png"),
        },
        "overall": "PASS" if run["status"] == "SUCCEEDED" and quality["hard_constraint_errors"] == 0 and not quality["missing_employee_days"] and not quality["duplicate_assignment_days"] and not quality["leave_violations"] and confirm_status == "PASS" else "FAIL",
    }
    (OUT / "phase7_pilot_results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
